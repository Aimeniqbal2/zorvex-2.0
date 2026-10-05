import logging
from rest_framework import viewsets, permissions, status, filters
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.exceptions import ValidationError
from django.core.exceptions import ValidationError as DjangoValidationError

logger = logging.getLogger(__name__)
from django.utils import timezone
from platform_core.permissions import ModulePermission
from erp_core.views import TenantModelViewSet
from erp_core.permissions import RolePermission
from .models import (
    SecurityProposal, SecurityServiceType, ClientLocation, SecurityAssessment,
    ProposalVersion, ProposalServiceLine, SecurityProposalMeeting,
    MeetingParticipant, ProposalFollowUp, SecurityProposalMeetingStatus,
    FollowUpStatus, SecurityProposalStatus, SecurityAssessmentStatus,
    AssessmentRiskFinding, AssessmentStaffingRecommendation,
    AssessmentEquipmentRecommendation, AssessmentAttachment, RiskLevel,
    ContractEquipmentRequirement, ProposalAdditionalCharge,
    ProposalSignedDocument
)
from .serializers import (
    SecurityProposalSerializer, SecurityServiceTypeSerializer, 
    ClientLocationSerializer, SecurityAssessmentSerializer,
    SecurityAssessmentDetailSerializer, ProposalVersionSerializer,
    ProposalVersionDetailSerializer, ProposalServiceLineSerializer,
    SecurityProposalMeetingSerializer, MeetingParticipantSerializer,
    ProposalFollowUpSerializer, AssessmentRiskFindingSerializer,
    AssessmentStaffingRecommendationSerializer,
    AssessmentEquipmentRecommendationSerializer,
    AssessmentAttachmentSerializer, ContractEquipmentRequirementSerializer,
    ProposalAdditionalChargeSerializer, ProposalSignedDocumentSerializer
)
from .services.workflow import SecurityProposalWorkflowService


class SecurityServiceTypeViewSet(TenantModelViewSet):
    queryset = SecurityServiceType.objects.all()
    serializer_class = SecurityServiceTypeSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']


class ClientLocationViewSet(TenantModelViewSet):
    queryset = ClientLocation.objects.all()
    serializer_class = ClientLocationSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['name', 'address', 'customer__name']
    ordering_fields = ['name', 'created_at']

    def get_queryset(self):
        qs = super().get_queryset()
        customer_id = self.request.query_params.get('customer') or self.request.query_params.get('customer_id')
        if customer_id:
            qs = qs.filter(customer_id=customer_id)
        is_active = self.request.query_params.get('is_active')
        if is_active is not None:
            if str(is_active).lower() in ['true', '1']:
                qs = qs.filter(is_active=True)
            elif str(is_active).lower() in ['false', '0']:
                qs = qs.filter(is_active=False)
        return qs

    def perform_create(self, serializer):
        loc = serializer.save()
        try:
            from operations.models import OperationalSite
            OperationalSite.objects.get_or_create(
                company=loc.company,
                crm_entity=loc.customer,
                name=loc.name,
                defaults={
                    'address': getattr(loc.crm_address, 'street_address', '') if loc.crm_address else (loc.notes or loc.name),
                    'is_active': loc.is_active
                }
            )
        except Exception as e:
            logger.warning(f"Could not auto-create OperationalSite for location {loc.name}: {e}")

    def perform_update(self, serializer):
        loc = serializer.save()
        try:
            from operations.models import OperationalSite
            OperationalSite.objects.filter(
                company=loc.company,
                crm_entity=loc.customer,
                name=loc.name
            ).update(is_active=loc.is_active)
        except Exception as e:
            logger.warning(f"Could not auto-update OperationalSite for location {loc.name}: {e}")



class SecurityProposalViewSet(TenantModelViewSet):
    queryset = SecurityProposal.objects.all()
    serializer_class = SecurityProposalSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['proposal_number', 'title', 'customer__name']
    ordering_fields = ['created_at', 'status', 'valid_until']

    def get_queryset(self):
        qs = super().get_queryset()
        customer_id = self.request.query_params.get('customer') or self.request.query_params.get('customer_id')
        if customer_id:
            qs = qs.filter(customer_id=customer_id)
        status_filter = self.request.query_params.get('status')
        if status_filter:
            qs = qs.filter(status=status_filter)
        return qs

    def perform_create(self, serializer):
        company = self.request.user.company
        proposal = serializer.save(
            company=company,
            status=SecurityProposalStatus.ACTIVE,
            is_handoff_ready=True
        )
        # Auto-create active version 1 for Final Requirements
        from django.db import transaction
        with transaction.atomic():
            ProposalVersion.objects.create(
                proposal=proposal,
                company=company,
                version_number=1,
                version_type='Final Requirement',
                status='ACTIVE'
            )
            try:
                SecurityProposalWorkflowService.sync_proposal_and_locations_to_operations(proposal)
            except Exception as e:
                logger.warning(f"Error auto-syncing new requirement to operations: {e}")

    @action(detail=True, methods=['post'], url_path='sync-to-operations')
    def sync_to_operations(self, request, pk=None):
        proposal = self.get_object()
        try:
            contract = SecurityProposalWorkflowService.sync_proposal_and_locations_to_operations(proposal)
            return Response({
                'success': True,
                'status': proposal.status,
                'contract_id': str(contract.id) if contract else None,
                'contract_code': contract.contract_code if contract else None,
                'message': 'Successfully synced final requirements into Operations!'
            }, status=status.HTTP_200_OK)
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'])
    def transition(self, request, pk=None):
        proposal = self.get_object()
        new_status = request.data.get('status')
        if not new_status:
            return Response({'error': 'Status is required'}, status=status.HTTP_400_BAD_REQUEST)
        
        try:
            SecurityProposalWorkflowService.transition_status(proposal, new_status, request.user)
            return Response({'status': proposal.status})
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'])
    def start_meeting_stage(self, request, pk=None):
        proposal = self.get_object()
        try:
            SecurityProposalWorkflowService.start_meeting_stage(proposal, user=request.user)
            return Response({
                'status': proposal.status,
                'message': 'Proposal transitioned to MEETING stage.'
            })
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'])
    def advance_to_site_assessment(self, request, pk=None):
        proposal = self.get_object()
        try:
            SecurityProposalWorkflowService.advance_to_site_assessment(proposal, user=request.user)
            return Response({
                'status': proposal.status,
                'message': 'Proposal advanced to SITE_ASSESSMENT stage.'
            })
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='advance-to-final-proposal')
    def advance_to_final_proposal(self, request, pk=None):
        proposal = self.get_object()
        try:
            SecurityProposalWorkflowService.advance_to_final_proposal(proposal, user=request.user)
            return Response({
                'status': proposal.status,
                'message': 'Proposal advanced to FINAL_PROPOSAL stage.'
            })
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['get'])
    def assessments(self, request, pk=None):
        proposal = self.get_object()
        assessments = proposal.assessments.filter(company=proposal.company).order_by('-created_at')
        serializer = SecurityAssessmentSerializer(assessments, many=True)
        return Response(serializer.data)

    @action(detail=True, methods=['get'])
    def meetings(self, request, pk=None):
        proposal = self.get_object()
        meetings = proposal.meetings.filter(company=proposal.company).order_by('-scheduled_at')
        serializer = SecurityProposalMeetingSerializer(meetings, many=True)
        return Response(serializer.data)

    @action(detail=True, methods=['get'])
    def follow_ups(self, request, pk=None):
        proposal = self.get_object()
        follow_ups = proposal.follow_ups.filter(company=proposal.company).order_by('status', 'due_at')
        serializer = ProposalFollowUpSerializer(follow_ups, many=True)
        return Response(serializer.data)

    @action(detail=True, methods=['post'])
    def schedule_meeting(self, request, pk=None):
        proposal = self.get_object()
        try:
            meeting = SecurityProposalWorkflowService.schedule_meeting(
                proposal=proposal,
                data=request.data.copy(),
                user=request.user
            )
            serializer = SecurityProposalMeetingSerializer(meeting)
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['get'])
    def prepare_email(self, request, pk=None):
        proposal = self.get_object()
        version_id = request.query_params.get('version_id')
        version = None
        if version_id:
            version = proposal.versions.filter(id=version_id, company=proposal.company).first()
        if not version:
            version = proposal.versions.filter(company=proposal.company).order_by('-version_number').first()

        from communications.models import EmailTemplate
        from django.template import Context, Template
        
        template = EmailTemplate.objects.filter(company=proposal.company, code='SECURITY_INITIAL_PROPOSAL').first()
        if not template:
            template = EmailTemplate(
                subject="Proposal {{ proposal.proposal_number }} - {{ company.name }}",
                body_text="Dear {{ customer.name }},\n\nThank you for the opportunity to submit our proposal.\n\nRegards,\n{{ company.name }}"
            )

        context_dict = {
            'company': proposal.company,
            'customer': proposal.customer,
            'proposal': proposal,
            'version': version,
            'user': request.user
        }
        
        try:
            subject_template = Template(template.subject)
            body_template = Template(template.body_text)
            ctx = Context(context_dict)
            rendered_subject = subject_template.render(ctx)
            rendered_body = body_template.render(ctx)
        except Exception as e:
            rendered_subject = template.subject
            rendered_body = template.body_text
            
        prefill_to = proposal.customer.email if hasattr(proposal.customer, 'email') else ''
        if not prefill_to:
            contact = proposal.customer.contacts.first() if hasattr(proposal.customer, 'contacts') else None
            if contact and hasattr(contact, 'email'):
                prefill_to = contact.email
                
        return Response({
            'subject': rendered_subject,
            'body': rendered_body,
            'to': prefill_to,
            'version_id': str(version.id) if version else None,
            'version_number': version.version_number if version else None
        })

    @action(detail=True, methods=['post'])
    def send_proposal_email(self, request, pk=None):
        proposal = self.get_object()
        email_id = request.data.get('email_id')
        version_id = request.data.get('version_id')
        version_number = request.data.get('version_number')
        
        if not email_id:
            return Response({'error': 'email_id is required'}, status=status.HTTP_400_BAD_REQUEST)
            
        from communications.models import OutboundEmail
        
        try:
            email = OutboundEmail.objects.get(id=email_id, company=proposal.company)
        except OutboundEmail.DoesNotExist:
            return Response({'error': 'Email not found'}, status=status.HTTP_404_NOT_FOUND)
            
        try:
            success, error, version = SecurityProposalWorkflowService.send_proposal_email(
                proposal=proposal,
                email=email,
                version_id=version_id,
                version_number=version_number,
                user=request.user
            )
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)
            
        email.refresh_from_db()
        proposal.refresh_from_db()
        
        if success:
            return Response({
                'status': 'sent',
                'proposal_status': proposal.status,
                'version_id': str(version.id) if version else None,
                'version_number': version.version_number if version else None,
                'is_frozen': version.is_frozen if version else None
            })
        else:
            return Response({'status': 'failed', 'error': error}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['get'])
    def versions(self, request, pk=None):
        proposal = self.get_object()
        versions = proposal.versions.filter(company=proposal.company).order_by('-version_number')
        serializer = ProposalVersionDetailSerializer(versions, many=True)
        return Response(serializer.data)

    @action(detail=True, methods=['post'], url_path='import-recommendations')
    def import_recommendations(self, request, pk=None):
        proposal = self.get_object()
        version_id = request.data.get('version_id')
        assessment_ids = request.data.get('assessment_ids')

        version = None
        if version_id:
            version = proposal.versions.filter(id=version_id, company=proposal.company).first()
        if not version:
            version = proposal.versions.filter(
                company=proposal.company,
                version_type__icontains='Final Proposal',
                is_frozen=False
            ).first()
        if not version:
            version = proposal.versions.filter(company=proposal.company, is_frozen=False).order_by('-version_number').first()

        if not version:
            return Response({'error': 'No active editable proposal version found. Please advance to Final Proposal or create a revision first.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            result = SecurityProposalWorkflowService.import_assessment_recommendations(
                proposal_version=version,
                user=request.user,
                assessment_ids=assessment_ids
            )
            version.refresh_from_db()
            return Response({
                'message': f"Imported {result['staffing_imported']} service lines and {result['equipment_imported']} equipment requirements.",
                'result': result,
                'version': ProposalVersionDetailSerializer(version).data
            })
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='create-final-revision')
    def create_final_revision(self, request, pk=None):
        proposal = self.get_object()
        base_version_id = request.data.get('base_version_id')

        try:
            new_version = SecurityProposalWorkflowService.create_final_proposal_revision(
                proposal=proposal,
                base_version_id=base_version_id,
                user=request.user
            )
            return Response({
                'message': f"Created revision v{new_version.version_number}.",
                'version': ProposalVersionDetailSerializer(new_version).data
            }, status=status.HTTP_201_CREATED)
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['get'], url_path='prepare-final-email')
    def prepare_final_email(self, request, pk=None):
        proposal = self.get_object()
        version_id = request.query_params.get('version_id')
        version = None
        if version_id:
            version = proposal.versions.filter(id=version_id, company=proposal.company).first()
        if not version:
            version = proposal.versions.filter(
                company=proposal.company,
                version_type__icontains='Final Proposal'
            ).order_by('-version_number').first()
        if not version:
            version = proposal.versions.filter(company=proposal.company).order_by('-version_number').first()

        from communications.models import EmailTemplate
        from django.template import Context, Template
        
        template = EmailTemplate.objects.filter(company=proposal.company, code='SECURITY_FINAL_PROPOSAL').first()
        if not template:
            template = EmailTemplate(
                subject="Final Proposal {{ proposal.proposal_number }} v{{ version.version_number }} - {{ company.name }}",
                body_text="Dear {{ customer.name }},\n\nWe are pleased to submit our comprehensive Final Security Proposal (Ref: {{ proposal.proposal_number }} v{{ version.version_number }}).\n\nSummary of Offer:\n- Monthly Recurring Total: PKR {{ version.total_monthly_recurring }}\n- One-Time Total: PKR {{ version.total_one_time }}\n- Grand Total: PKR {{ version.grand_total }}\n\nPlease review the attached commercial terms and requirements.\n\nRegards,\n{{ company.name }}"
            )

        context_dict = {
            'company': proposal.company,
            'customer': proposal.customer,
            'proposal': proposal,
            'version': version,
            'user': request.user
        }
        
        try:
            subject_template = Template(template.subject)
            body_template = Template(template.body_text)
            ctx = Context(context_dict)
            rendered_subject = subject_template.render(ctx)
            rendered_body = body_template.render(ctx)
        except Exception:
            rendered_subject = template.subject
            rendered_body = template.body_text
            
        prefill_to = proposal.customer.email if hasattr(proposal.customer, 'email') else ''
        if not prefill_to:
            contact = proposal.customer.contacts.first() if hasattr(proposal.customer, 'contacts') else None
            if contact and hasattr(contact, 'email'):
                prefill_to = contact.email
                
        return Response({
            'subject': rendered_subject,
            'body': rendered_body,
            'to': prefill_to,
            'version_id': str(version.id) if version else None,
            'version_number': version.version_number if version else None
        })

    @action(detail=True, methods=['post'], url_path='send-final-proposal-email')
    def send_final_proposal_email(self, request, pk=None):
        proposal = self.get_object()
        email_id = request.data.get('email_id')
        version_id = request.data.get('version_id')
        version_number = request.data.get('version_number')
        
        if not email_id:
            return Response({'error': 'email_id is required'}, status=status.HTTP_400_BAD_REQUEST)
            
        from communications.models import OutboundEmail
        
        try:
            email = OutboundEmail.objects.get(id=email_id, company=proposal.company)
        except OutboundEmail.DoesNotExist:
            return Response({'error': 'Email not found'}, status=status.HTTP_404_NOT_FOUND)
            
        try:
            success, error, version = SecurityProposalWorkflowService.send_final_proposal_email(
                proposal=proposal,
                email=email,
                version_id=version_id,
                version_number=version_number,
                user=request.user
            )
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)
            
        email.refresh_from_db()
        proposal.refresh_from_db()
        
        if success:
            return Response({
                'status': 'sent',
                'proposal_status': proposal.status,
                'version_id': str(version.id) if version else None,
                'version_number': version.version_number if version else None,
                'is_frozen': version.is_frozen if version else None
            })
        else:
            return Response({'status': 'failed', 'error': error}, status=status.HTTP_400_BAD_REQUEST)

    # -------------------------------------------------------------
    # PHASE S-2G: APPROVAL, SIGNING & ACTIVATION ACTIONS
    # -------------------------------------------------------------

    @action(detail=True, methods=['post'])
    def approve(self, request, pk=None):
        proposal = self.get_object()
        try:
            proposal = SecurityProposalWorkflowService.approve_proposal(
                proposal=proposal,
                data=request.data,
                user=request.user
            )
            return Response(SecurityProposalSerializer(proposal).data)
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'])
    def reject(self, request, pk=None):
        proposal = self.get_object()
        try:
            proposal = SecurityProposalWorkflowService.reject_proposal(
                proposal=proposal,
                data=request.data,
                user=request.user
            )
            return Response(SecurityProposalSerializer(proposal).data)
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='put-on-hold')
    def put_on_hold(self, request, pk=None):
        proposal = self.get_object()
        try:
            proposal = SecurityProposalWorkflowService.put_proposal_on_hold(
                proposal=proposal,
                data=request.data,
                user=request.user
            )
            return Response(SecurityProposalSerializer(proposal).data)
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='resume-from-on-hold')
    def resume_from_on_hold(self, request, pk=None):
        proposal = self.get_object()
        target_status = request.data.get('target_status')
        try:
            proposal = SecurityProposalWorkflowService.resume_proposal_from_on_hold(
                proposal=proposal,
                target_status=target_status,
                user=request.user
            )
            return Response(SecurityProposalSerializer(proposal).data)
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='start-signing')
    def start_signing(self, request, pk=None):
        proposal = self.get_object()
        try:
            proposal = SecurityProposalWorkflowService.start_signing(
                proposal=proposal,
                data=request.data,
                user=request.user
            )
            return Response(SecurityProposalSerializer(proposal).data)
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='complete-signing')
    def complete_signing(self, request, pk=None):
        proposal = self.get_object()
        try:
            proposal = SecurityProposalWorkflowService.complete_signing(
                proposal=proposal,
                data=request.data,
                user=request.user
            )
            return Response(SecurityProposalSerializer(proposal).data)
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='upload-signed-document')
    def upload_signed_document(self, request, pk=None):
        proposal = self.get_object()
        file_obj = request.FILES.get('file')
        title = request.data.get('title', '').strip()
        document_type = request.data.get('document_type', 'SIGNED_CONTRACT')
        notes = request.data.get('notes', '').strip()

        if not title and file_obj:
            title = file_obj.name

        if not title:
            return Response({'error': 'Title is required for the document.'}, status=status.HTTP_400_BAD_REQUEST)

        doc = ProposalSignedDocument.objects.create(
            company=proposal.company,
            proposal=proposal,
            title=title,
            document_type=document_type,
            file=file_obj,
            notes=notes,
            uploaded_by=request.user
        )

        return Response(ProposalSignedDocumentSerializer(doc).data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'], url_path='activate-client')
    def activate_client(self, request, pk=None):
        proposal = self.get_object()
        try:
            proposal = SecurityProposalWorkflowService.activate_client(
                proposal=proposal,
                user=request.user
            )
            return Response({
                'message': 'Client activated successfully.',
                'proposal': SecurityProposalSerializer(proposal).data,
                'contract_id': str(proposal.contract.id) if proposal.contract else None,
                'contract_code': proposal.contract.contract_code if proposal.contract else None,
                'sites_count': proposal.contract.sites.count() if proposal.contract else 0
            })
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['get'], url_path='handoff-summary')
    def handoff_summary(self, request, pk=None):
        """
        Retrieves the structured cross-module handoff summary for an active proposal.
        """
        from security_crm.services.handoff import CRMCrossModuleHandoffService
        proposal = self.get_object()
        summary = CRMCrossModuleHandoffService.get_handoff_summary(proposal)
        return Response(summary, status=status.HTTP_200_OK)

    @action(detail=True, methods=['post'], url_path='prepare-handoff')
    def prepare_handoff(self, request, pk=None):
        """
        Prepares and locks the cross-module handoff payload for downstream security modules.
        """
        from security_crm.services.handoff import CRMCrossModuleHandoffService
        proposal = self.get_object()
        notes = request.data.get('notes', '')
        try:
            summary = CRMCrossModuleHandoffService.prepare_cross_module_handoff(
                proposal=proposal,
                user=request.user,
                notes=notes
            )
            return Response({
                'message': 'Cross-module handoff prepared and certified successfully.',
                'handoff_summary': summary
            }, status=status.HTTP_200_OK)
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)


class ProposalSignedDocumentViewSet(TenantModelViewSet):
    queryset = ProposalSignedDocument.objects.all()
    serializer_class = ProposalSignedDocumentSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']
    filter_backends = [filters.SearchFilter]
    search_fields = ['title', 'notes']

    def get_queryset(self):
        qs = super().get_queryset()
        proposal_id = self.request.query_params.get('proposal') or self.request.query_params.get('proposal_id')
        if proposal_id:
            qs = qs.filter(proposal_id=proposal_id)
        doc_type = self.request.query_params.get('document_type')
        if doc_type:
            qs = qs.filter(document_type=doc_type)
        return qs

    def perform_create(self, serializer):
        company = self.request.user.company
        proposal = serializer.validated_data.get('proposal')
        if proposal and str(proposal.company_id) != str(company.id):
            raise ValidationError({'proposal': 'Proposal must belong to the same company.'})
        serializer.save(company=company, uploaded_by=self.request.user)


class SecurityProposalMeetingViewSet(TenantModelViewSet):
    queryset = SecurityProposalMeeting.objects.all()
    serializer_class = SecurityProposalMeetingSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['subject', 'location', 'agenda', 'discussion_notes']
    ordering_fields = ['scheduled_at', 'status', 'created_at']

    def get_queryset(self):
        qs = super().get_queryset()
        proposal_id = self.request.query_params.get('proposal') or self.request.query_params.get('proposal_id')
        if proposal_id:
            qs = qs.filter(proposal_id=proposal_id)
        meeting_status = self.request.query_params.get('status')
        if meeting_status:
            qs = qs.filter(status=meeting_status)
        meeting_type = self.request.query_params.get('meeting_type')
        if meeting_type:
            qs = qs.filter(meeting_type=meeting_type)
        outcome = self.request.query_params.get('outcome')
        if outcome:
            qs = qs.filter(outcome=outcome)
        return qs

    def perform_create(self, serializer):
        company = self.request.user.company
        proposal = serializer.validated_data.get('proposal')
        
        # Validate proposal tenant
        if proposal and str(proposal.company_id) != str(company.id):
            raise ValidationError({'proposal': 'Proposal must belong to the same company.'})

        # Auto-advance proposal to MEETING if it is currently in SENT
        if proposal and proposal.status == SecurityProposalStatus.SENT:
            SecurityProposalWorkflowService.start_meeting_stage(proposal, user=self.request.user)

        serializer.save(company=company, created_by=self.request.user)

    @action(detail=True, methods=['post'])
    def complete(self, request, pk=None):
        meeting = self.get_object()
        outcome = request.data.get('outcome', meeting.outcome)
        outcome_notes = request.data.get('outcome_notes', meeting.outcome_notes)
        discussion_notes = request.data.get('discussion_notes', meeting.discussion_notes)
        client_requirements = request.data.get('client_requirements', meeting.client_requirements)
        commercial_concerns = request.data.get('commercial_concerns', meeting.commercial_concerns)
        operational_concerns = request.data.get('operational_concerns', meeting.operational_concerns)
        agreed_points = request.data.get('agreed_points', meeting.agreed_points)
        pending_items = request.data.get('pending_items', meeting.pending_items)

        meeting.status = SecurityProposalMeetingStatus.COMPLETED
        meeting.outcome = outcome
        meeting.outcome_notes = outcome_notes
        meeting.discussion_notes = discussion_notes
        meeting.client_requirements = client_requirements
        meeting.commercial_concerns = commercial_concerns
        meeting.operational_concerns = operational_concerns
        meeting.agreed_points = agreed_points
        meeting.pending_items = pending_items
        meeting.completed_at = timezone.now()
        meeting.save()

        # Optional follow-up creation
        follow_up_title = request.data.get('follow_up_title')
        if follow_up_title:
            ProposalFollowUp.objects.create(
                company=meeting.company,
                proposal=meeting.proposal,
                related_meeting=meeting,
                title=follow_up_title,
                description=request.data.get('follow_up_description', ''),
                due_at=request.data.get('follow_up_due_at'),
                priority=request.data.get('follow_up_priority', 'MEDIUM'),
                assigned_to_id=request.data.get('follow_up_assigned_to'),
                created_by=request.user
            )

        meeting.refresh_from_db()
        return Response(SecurityProposalMeetingSerializer(meeting).data)

    @action(detail=True, methods=['post'])
    def cancel(self, request, pk=None):
        meeting = self.get_object()
        meeting.status = SecurityProposalMeetingStatus.CANCELLED
        meeting.save(update_fields=['status', 'updated_at'])
        return Response(SecurityProposalMeetingSerializer(meeting).data)

    @action(detail=True, methods=['post'])
    def no_show(self, request, pk=None):
        meeting = self.get_object()
        meeting.status = SecurityProposalMeetingStatus.NO_SHOW
        meeting.save(update_fields=['status', 'updated_at'])
        return Response(SecurityProposalMeetingSerializer(meeting).data)

    @action(detail=True, methods=['post'])
    def add_participant(self, request, pk=None):
        meeting = self.get_object()
        data = request.data.copy()
        data['meeting'] = meeting.id
        serializer = MeetingParticipantSerializer(data=data)
        serializer.is_valid(raise_exception=True)
        participant = serializer.save(company=meeting.company)
        return Response(MeetingParticipantSerializer(participant).data, status=status.HTTP_201_CREATED)


    @action(detail=True, methods=['post'])
    def add_follow_up(self, request, pk=None):
        meeting = self.get_object()
        data = request.data.copy()
        data['proposal'] = meeting.proposal.id
        data['related_meeting'] = meeting.id
        serializer = ProposalFollowUpSerializer(data=data)
        serializer.is_valid(raise_exception=True)
        follow_up = serializer.save(company=meeting.company, created_by=request.user)
        return Response(ProposalFollowUpSerializer(follow_up).data, status=status.HTTP_201_CREATED)


class MeetingParticipantViewSet(TenantModelViewSet):
    queryset = MeetingParticipant.objects.all()
    serializer_class = MeetingParticipantSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']

    def get_queryset(self):
        qs = super().get_queryset()
        meeting_id = self.request.query_params.get('meeting') or self.request.query_params.get('meeting_id')
        if meeting_id:
            qs = qs.filter(meeting_id=meeting_id)
        participant_type = self.request.query_params.get('participant_type')
        if participant_type:
            qs = qs.filter(participant_type=participant_type)
        return qs

    def perform_create(self, serializer):
        company = self.request.user.company
        meeting = serializer.validated_data.get('meeting')
        crm_contact = serializer.validated_data.get('crm_contact')
        user = serializer.validated_data.get('user')

        if meeting and str(meeting.company_id) != str(company.id):
            raise ValidationError({'meeting': 'Meeting must belong to the same company.'})
        if crm_contact:
            if str(crm_contact.company_id) != str(company.id):
                raise ValidationError({'crm_contact': 'Contact must belong to the same company.'})
            if meeting and str(crm_contact.entity_id) != str(meeting.proposal.customer_id):
                raise ValidationError({'crm_contact': 'Contact must belong to the proposal customer.'})
        if user and str(user.company_id) != str(company.id):
            raise ValidationError({'user': 'User must belong to the same company.'})

        serializer.save(company=company)


class ProposalFollowUpViewSet(TenantModelViewSet):
    queryset = ProposalFollowUp.objects.all()
    serializer_class = ProposalFollowUpSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['title', 'description']
    ordering_fields = ['due_at', 'status', 'priority', 'created_at']

    def get_queryset(self):
        qs = super().get_queryset()
        proposal_id = self.request.query_params.get('proposal') or self.request.query_params.get('proposal_id')
        if proposal_id:
            qs = qs.filter(proposal_id=proposal_id)
        related_meeting = self.request.query_params.get('related_meeting') or self.request.query_params.get('related_meeting_id')
        if related_meeting:
            qs = qs.filter(related_meeting_id=related_meeting)
        st = self.request.query_params.get('status')
        if st:
            qs = qs.filter(status=st)
        priority = self.request.query_params.get('priority')
        if priority:
            qs = qs.filter(priority=priority)
        assigned_to = self.request.query_params.get('assigned_to')
        if assigned_to:
            qs = qs.filter(assigned_to_id=assigned_to)
        return qs

    def perform_create(self, serializer):
        company = self.request.user.company
        proposal = serializer.validated_data.get('proposal')
        related_meeting = serializer.validated_data.get('related_meeting')
        assigned_to = serializer.validated_data.get('assigned_to')

        if proposal and str(proposal.company_id) != str(company.id):
            raise ValidationError({'proposal': 'Proposal must belong to the same company.'})
        if related_meeting:
            if str(related_meeting.company_id) != str(company.id) or str(related_meeting.proposal_id) != str(proposal.id):
                raise ValidationError({'related_meeting': 'Related meeting must belong to this proposal and company.'})
        if assigned_to and str(assigned_to.company_id) != str(company.id):
            raise ValidationError({'assigned_to': 'Assigned user must belong to the same company.'})

        serializer.save(company=company, created_by=self.request.user)

    @action(detail=True, methods=['post'])
    def complete(self, request, pk=None):
        follow_up = self.get_object()
        follow_up.status = FollowUpStatus.COMPLETED
        follow_up.completed_at = timezone.now()
        follow_up.save(update_fields=['status', 'completed_at', 'updated_at'])
        return Response(ProposalFollowUpSerializer(follow_up).data)

    @action(detail=True, methods=['post'])
    def cancel(self, request, pk=None):
        follow_up = self.get_object()
        follow_up.status = FollowUpStatus.CANCELLED
        follow_up.save(update_fields=['status', 'updated_at'])
        return Response(ProposalFollowUpSerializer(follow_up).data)


class SecurityAssessmentViewSet(TenantModelViewSet):
    queryset = SecurityAssessment.objects.all()
    serializer_class = SecurityAssessmentSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']
    filter_backends = [filters.SearchFilter]
    search_fields = ['client_location__name', 'proposal__proposal_number', 'site_overview']

    def get_queryset(self):
        qs = super().get_queryset()
        proposal_id = self.request.query_params.get('proposal') or self.request.query_params.get('proposal_id')
        if proposal_id:
            qs = qs.filter(proposal_id=proposal_id)
        location_id = self.request.query_params.get('client_location') or self.request.query_params.get('location_id')
        if location_id:
            qs = qs.filter(client_location_id=location_id)
        status_filter = self.request.query_params.get('status')
        if status_filter:
            qs = qs.filter(status=status_filter)
        customer_id = self.request.query_params.get('customer') or self.request.query_params.get('customer_id')
        if customer_id:
            from django.db.models import Q
            qs = qs.filter(Q(proposal__customer_id=customer_id) | Q(client_location__customer_id=customer_id))
        return qs

    def get_serializer_class(self):
        if self.action in ['retrieve', 'complete', 'reopen']:
            return SecurityAssessmentDetailSerializer
        return SecurityAssessmentSerializer

    def perform_create(self, serializer):
        company = self.request.user.company
        proposal = serializer.validated_data.get('proposal')
        client_location = serializer.validated_data.get('client_location')

        if proposal and str(proposal.company_id) != str(company.id):
            raise ValidationError({'proposal': 'Proposal must belong to the same company.'})
        if client_location and str(client_location.company_id) != str(company.id):
            raise ValidationError({'client_location': 'Client Location must belong to the same company.'})
        if client_location and proposal and str(client_location.customer_id) != str(proposal.customer_id):
            raise ValidationError({'client_location': 'Client location does not belong to the proposal customer.'})

        assessed_by = serializer.validated_data.get('assessed_by') or self.request.user
        serializer.save(company=company, assessed_by=assessed_by)

    def perform_update(self, serializer):
        instance = self.get_object()
        if instance.status == SecurityAssessmentStatus.COMPLETED:
            new_status = serializer.validated_data.get('status')
            if new_status == SecurityAssessmentStatus.COMPLETED or not new_status:
                raise ValidationError({'detail': 'Completed assessments are locked against modification. Reopen assessment first.'})
        serializer.save()

    @action(detail=True, methods=['post'])
    def complete(self, request, pk=None):
        assessment = self.get_object()
        assessment.status = SecurityAssessmentStatus.COMPLETED
        assessment.completed_at = timezone.now()
        assessment.completed_by = request.user
        assessment.save(update_fields=['status', 'completed_at', 'completed_by', 'updated_at'])
        return Response(SecurityAssessmentDetailSerializer(assessment).data)

    @action(detail=True, methods=['post'])
    def reopen(self, request, pk=None):
        assessment = self.get_object()
        assessment.status = SecurityAssessmentStatus.IN_PROGRESS
        assessment.save(update_fields=['status', 'updated_at'])
        return Response(SecurityAssessmentDetailSerializer(assessment).data)

    @action(detail=True, methods=['post'], url_path='add-risk')
    def add_risk(self, request, pk=None):
        assessment = self.get_object()
        if assessment.status == SecurityAssessmentStatus.COMPLETED:
            return Response({'error': 'Cannot add risk findings to a completed assessment.'}, status=status.HTTP_400_BAD_REQUEST)
        
        serializer = AssessmentRiskFindingSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save(company=assessment.company, assessment=assessment)
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='add-staffing')
    def add_staffing(self, request, pk=None):
        assessment = self.get_object()
        if assessment.status == SecurityAssessmentStatus.COMPLETED:
            return Response({'error': 'Cannot add staffing recommendations to a completed assessment.'}, status=status.HTTP_400_BAD_REQUEST)
        
        serializer = AssessmentStaffingRecommendationSerializer(data=request.data)
        if serializer.is_valid():
            service_type = serializer.validated_data.get('service_type')
            if service_type and str(service_type.company_id) != str(assessment.company_id):
                return Response({'error': 'ServiceType company mismatch.'}, status=status.HTTP_400_BAD_REQUEST)
            serializer.save(company=assessment.company, assessment=assessment)
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='add-equipment')
    def add_equipment(self, request, pk=None):
        assessment = self.get_object()
        if assessment.status == SecurityAssessmentStatus.COMPLETED:
            return Response({'error': 'Cannot add equipment recommendations to a completed assessment.'}, status=status.HTTP_400_BAD_REQUEST)
        
        serializer = AssessmentEquipmentRecommendationSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save(company=assessment.company, assessment=assessment)
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='add-attachment')
    def add_attachment(self, request, pk=None):
        assessment = self.get_object()
        if assessment.status == SecurityAssessmentStatus.COMPLETED:
            return Response({'error': 'Cannot add attachments to a completed assessment.'}, status=status.HTTP_400_BAD_REQUEST)
        
        serializer = AssessmentAttachmentSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save(company=assessment.company, assessment=assessment, uploaded_by=request.user)
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class AssessmentRiskFindingViewSet(TenantModelViewSet):
    queryset = AssessmentRiskFinding.objects.all()
    serializer_class = AssessmentRiskFindingSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']
    filter_backends = [filters.SearchFilter]
    search_fields = ['hazard_title', 'description', 'control_measures']

    def get_queryset(self):
        qs = super().get_queryset()
        assessment_id = self.request.query_params.get('assessment') or self.request.query_params.get('assessment_id')
        if assessment_id:
            qs = qs.filter(assessment_id=assessment_id)
        risk_level = self.request.query_params.get('risk_level')
        if risk_level:
            qs = qs.filter(risk_level=risk_level)
        status_filter = self.request.query_params.get('status')
        if status_filter:
            qs = qs.filter(status=status_filter)
        return qs

    def perform_create(self, serializer):
        assessment = serializer.validated_data.get('assessment')
        if assessment and str(assessment.company_id) != str(self.request.user.company_id):
            raise ValidationError({'assessment': 'Assessment company mismatch.'})
        serializer.save(company=self.request.user.company)


class AssessmentStaffingRecommendationViewSet(TenantModelViewSet):
    queryset = AssessmentStaffingRecommendation.objects.all()
    serializer_class = AssessmentStaffingRecommendationSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']

    def get_queryset(self):
        qs = super().get_queryset()
        assessment_id = self.request.query_params.get('assessment') or self.request.query_params.get('assessment_id')
        if assessment_id:
            qs = qs.filter(assessment_id=assessment_id)
        service_type = self.request.query_params.get('service_type')
        if service_type:
            qs = qs.filter(service_type_id=service_type)
        return qs

    def perform_create(self, serializer):
        assessment = serializer.validated_data.get('assessment')
        service_type = serializer.validated_data.get('service_type')
        if assessment and str(assessment.company_id) != str(self.request.user.company_id):
            raise ValidationError({'assessment': 'Assessment company mismatch.'})
        if service_type and str(service_type.company_id) != str(self.request.user.company_id):
            raise ValidationError({'service_type': 'ServiceType company mismatch.'})
        serializer.save(company=self.request.user.company)


class AssessmentEquipmentRecommendationViewSet(TenantModelViewSet):
    queryset = AssessmentEquipmentRecommendation.objects.all()
    serializer_class = AssessmentEquipmentRecommendationSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']

    def get_queryset(self):
        qs = super().get_queryset()
        assessment_id = self.request.query_params.get('assessment') or self.request.query_params.get('assessment_id')
        if assessment_id:
            qs = qs.filter(assessment_id=assessment_id)
        return qs

    def perform_create(self, serializer):
        assessment = serializer.validated_data.get('assessment')
        if assessment and str(assessment.company_id) != str(self.request.user.company_id):
            raise ValidationError({'assessment': 'Assessment company mismatch.'})
        serializer.save(company=self.request.user.company)


class AssessmentAttachmentViewSet(TenantModelViewSet):
    queryset = AssessmentAttachment.objects.all()
    serializer_class = AssessmentAttachmentSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']

    def get_queryset(self):
        qs = super().get_queryset()
        assessment_id = self.request.query_params.get('assessment') or self.request.query_params.get('assessment_id')
        if assessment_id:
            qs = qs.filter(assessment_id=assessment_id)
        category = self.request.query_params.get('category')
        if category:
            qs = qs.filter(category=category)
        return qs

    def perform_create(self, serializer):
        assessment = serializer.validated_data.get('assessment')
        if assessment and str(assessment.company_id) != str(self.request.user.company_id):
            raise ValidationError({'assessment': 'Assessment company mismatch.'})
        serializer.save(company=self.request.user.company, uploaded_by=self.request.user)



class ProposalVersionViewSet(TenantModelViewSet):
    queryset = ProposalVersion.objects.all()
    serializer_class = ProposalVersionSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']

    def get_queryset(self):
        qs = super().get_queryset()
        proposal_id = self.request.query_params.get('proposal') or self.request.query_params.get('proposal_id')
        if proposal_id:
            qs = qs.filter(proposal_id=proposal_id)
        version_number = self.request.query_params.get('version_number')
        if version_number:
            qs = qs.filter(version_number=version_number)
        is_frozen = self.request.query_params.get('is_frozen')
        if is_frozen is not None:
            if str(is_frozen).lower() in ['true', '1']:
                qs = qs.filter(is_frozen=True)
            elif str(is_frozen).lower() in ['false', '0']:
                qs = qs.filter(is_frozen=False)
        return qs

    def get_serializer_class(self):
        if self.action in ['retrieve', 'import_recommendations', 'create_revision']:
            return ProposalVersionDetailSerializer
        return ProposalVersionSerializer

    def perform_create(self, serializer):
        proposal = serializer.validated_data.get('proposal')
        if proposal and str(proposal.company_id) != str(self.request.user.company_id):
            raise ValidationError({'proposal': 'Proposal company mismatch.'})
        serializer.save(company=self.request.user.company)

    def perform_update(self, serializer):
        instance = serializer.save()
        try:
            if instance.proposal:
                SecurityProposalWorkflowService.sync_proposal_and_locations_to_operations(instance.proposal)
        except Exception as e:
            logger.warning(f"Error auto-syncing version update to operations: {e}")

    @action(detail=True, methods=['post'], url_path='import-recommendations')
    def import_recommendations(self, request, pk=None):
        version = self.get_object()
        assessment_ids = request.data.get('assessment_ids')
        try:
            result = SecurityProposalWorkflowService.import_assessment_recommendations(
                proposal_version=version,
                user=request.user,
                assessment_ids=assessment_ids
            )
            version.refresh_from_db()
            return Response({
                'message': f"Imported {result['staffing_imported']} service lines and {result['equipment_imported']} equipment requirements.",
                'result': result,
                'version': ProposalVersionDetailSerializer(version).data
            })
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='create-revision')
    def create_revision(self, request, pk=None):
        version = self.get_object()
        try:
            new_version = SecurityProposalWorkflowService.create_final_proposal_revision(
                proposal=version.proposal,
                base_version_id=version.id,
                user=request.user
            )
            return Response({
                'message': f"Created revision v{new_version.version_number}.",
                'version': ProposalVersionDetailSerializer(new_version).data
            }, status=status.HTTP_201_CREATED)
        except (ValidationError, DjangoValidationError) as ve:
            return Response({'error': str(ve)}, status=status.HTTP_400_BAD_REQUEST)


class ProposalServiceLineViewSet(TenantModelViewSet):
    queryset = ProposalServiceLine.objects.all()
    serializer_class = ProposalServiceLineSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']

    def get_queryset(self):
        qs = super().get_queryset()
        version_id = self.request.query_params.get('proposal_version') or self.request.query_params.get('proposal_version_id')
        if version_id:
            qs = qs.filter(proposal_version_id=version_id)
        location_id = self.request.query_params.get('location') or self.request.query_params.get('location_id')
        if location_id:
            qs = qs.filter(location_id=location_id)
        service_type = self.request.query_params.get('service_type')
        if service_type:
            qs = qs.filter(service_type_id=service_type)
        return qs

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        version = serializer.validated_data.get('proposal_version')
        location = serializer.validated_data.get('location')
        service_type = serializer.validated_data.get('service_type')

        if version and version.is_frozen:
            raise ValidationError({'detail': 'Cannot add service lines to a frozen proposal version.'})

        # Check for existing line to upgrade instead of creating duplicate
        existing_line = ProposalServiceLine.objects.filter(
            proposal_version=version,
            location=location,
            service_type=service_type
        ).first()

        if not existing_line and service_type:
            st_name = (service_type.name or '').lower()
            if st_name in ['security guard', 'security guard civil', 'security guard (civil)']:
                existing_line = ProposalServiceLine.objects.filter(
                    proposal_version=version,
                    location=location,
                    service_type__name__in=['Security Guard', 'Security Guard (Civil)', 'Security Guard Civil']
                ).first()

        if existing_line:
            # Upgrade existing line
            for field, val in serializer.validated_data.items():
                if field not in ['proposal_version', 'location']:
                    setattr(existing_line, field, val)
            existing_line.save()
            try:
                SecurityProposalWorkflowService.sync_proposal_and_locations_to_operations(version.proposal)
            except Exception as e:
                logger.warning(f"Error auto-syncing updated service line to operations: {e}")
            return Response(self.get_serializer(existing_line).data, status=status.HTTP_200_OK)

        self.perform_create(serializer)
        headers = self.get_success_headers(serializer.data)
        return Response(serializer.data, status=status.HTTP_201_CREATED, headers=headers)

    def perform_create(self, serializer):
        version = serializer.validated_data.get('proposal_version')
        location = serializer.validated_data.get('location')
        service_type = serializer.validated_data.get('service_type')

        company = getattr(self.request.user, 'company', None)
        if not company:
            from erp_core.middleware import get_current_company
            from companies.models import Company
            cid = get_current_company() or self.request.META.get('HTTP_X_COMPANY_ID') or getattr(self.request.user, 'company_id', None)
            company = Company.objects.filter(id=cid).first() if cid else Company.objects.first()

        if version and str(version.company_id) != str(company.id):
            raise ValidationError({'proposal_version': 'ProposalVersion company mismatch.'})
        if location and str(location.company_id) != str(company.id):
            raise ValidationError({'location': 'Location company mismatch.'})
        if service_type and str(service_type.company_id) != str(company.id):
            raise ValidationError({'service_type': 'ServiceType company mismatch.'})

        instance = serializer.save(company=company)
        try:
            if instance.proposal_version and instance.proposal_version.proposal:
                SecurityProposalWorkflowService.sync_proposal_and_locations_to_operations(instance.proposal_version.proposal)
        except Exception as e:
            logger.warning(f"Error auto-syncing created service line to operations: {e}")

    def perform_update(self, serializer):
        instance = serializer.save()
        try:
            if instance.proposal_version and instance.proposal_version.proposal:
                SecurityProposalWorkflowService.sync_proposal_and_locations_to_operations(instance.proposal_version.proposal)
        except Exception as e:
            logger.warning(f"Error auto-syncing updated service line to operations: {e}")

    def perform_destroy(self, instance):
        proposal = instance.proposal_version.proposal if instance.proposal_version else None
        super().perform_destroy(instance)
        if proposal:
            try:
                SecurityProposalWorkflowService.sync_proposal_and_locations_to_operations(proposal)
            except Exception as e:
                logger.warning(f"Error auto-syncing deleted service line to operations: {e}")


class ContractEquipmentRequirementViewSet(TenantModelViewSet):
    queryset = ContractEquipmentRequirement.objects.all()
    serializer_class = ContractEquipmentRequirementSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']

    def get_queryset(self):
        qs = super().get_queryset()
        version_id = self.request.query_params.get('proposal_version') or self.request.query_params.get('proposal_version_id')
        if version_id:
            qs = qs.filter(proposal_version_id=version_id)
        location_id = self.request.query_params.get('location') or self.request.query_params.get('location_id')
        if location_id:
            qs = qs.filter(location_id=location_id)
        charge_type = self.request.query_params.get('charge_type')
        if charge_type:
            qs = qs.filter(charge_type=charge_type)
        return qs

    def perform_create(self, serializer):
        version = serializer.validated_data.get('proposal_version')
        location = serializer.validated_data.get('location')
        inventory_item = serializer.validated_data.get('inventory_item')

        if version and str(version.company_id) != str(self.request.user.company_id):
            raise ValidationError({'proposal_version': 'ProposalVersion company mismatch.'})
        if version and version.is_frozen:
            raise ValidationError({'detail': 'Cannot add equipment to a frozen proposal version.'})
        if location and str(location.company_id) != str(self.request.user.company_id):
            raise ValidationError({'location': 'Location company mismatch.'})
        if inventory_item and str(inventory_item.company_id) != str(self.request.user.company_id):
            raise ValidationError({'inventory_item': 'Inventory item company mismatch.'})

        serializer.save(company=self.request.user.company)

    def perform_update(self, serializer):
        instance = self.get_object()
        if instance.proposal_version.is_frozen:
            raise ValidationError({'detail': 'Cannot modify equipment in a frozen proposal version.'})
        serializer.save()

    def perform_destroy(self, instance):
        if instance.proposal_version.is_frozen:
            raise ValidationError({'detail': 'Cannot delete equipment from a frozen proposal version.'})
        super().perform_destroy(instance)


class ProposalAdditionalChargeViewSet(TenantModelViewSet):
    queryset = ProposalAdditionalCharge.objects.all()
    serializer_class = ProposalAdditionalChargeSerializer
    permission_classes = [permissions.IsAuthenticated, RolePermission, ModulePermission]
    required_module = 'crm'
    allowed_roles = ['admin', 'manager', 'sales']
    allowed_reads = ['admin', 'manager', 'sales', 'staff']

    def get_queryset(self):
        qs = super().get_queryset()
        version_id = self.request.query_params.get('proposal_version') or self.request.query_params.get('proposal_version_id')
        if version_id:
            qs = qs.filter(proposal_version_id=version_id)
        charge_type = self.request.query_params.get('charge_type')
        if charge_type:
            qs = qs.filter(charge_type=charge_type)
        return qs

    def perform_create(self, serializer):
        version = serializer.validated_data.get('proposal_version')
        if version and str(version.company_id) != str(self.request.user.company_id):
            raise ValidationError({'proposal_version': 'ProposalVersion company mismatch.'})
        if version and version.is_frozen:
            raise ValidationError({'detail': 'Cannot add charges to a frozen proposal version.'})

        serializer.save(company=self.request.user.company)

    def perform_update(self, serializer):
        instance = self.get_object()
        if instance.proposal_version.is_frozen:
            raise ValidationError({'detail': 'Cannot modify charges in a frozen proposal version.'})
        serializer.save()

    def perform_destroy(self, instance):
        if instance.proposal_version.is_frozen:
            raise ValidationError({'detail': 'Cannot delete charges from a frozen proposal version.'})
        super().perform_destroy(instance)


# ============================================================================
# FAST COSTING GRID & SHEET 1 EXCEL IMPORTER VIEWS
# ============================================================================

from rest_framework.views import APIView
from decimal import Decimal
from django.db import transaction
from crm.models import CRMEntity
import openpyxl


class CostingGridView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        from erp_core.middleware import get_current_company
        from companies.models import Company

        company_id = get_current_company() or request.META.get('HTTP_X_COMPANY_ID') or getattr(request.user, 'company_id', None)
        company = Company.objects.filter(id=company_id).first() if company_id else getattr(request.user, 'company', None)
        if not company and request.user.is_superuser:
            company = Company.objects.first()
        if not company:
            return Response([], status=status.HTTP_200_OK)

        versions = ProposalVersion.objects.filter(
            company=company,
            proposal__customer__isnull=False
        ).select_related(
            'proposal', 'proposal__customer'
        ).prefetch_related(
            'service_lines', 'service_lines__location', 'service_lines__service_type'
        ).order_by('-proposal__created_at', '-version_number')[:1000]

        data = []
        seen_customer_locations = set()
        for ver in versions:
            prop = ver.proposal
            cust = prop.customer
            lines = list(ver.service_lines.all())
            if not lines:
                continue

            loc_map = {}
            for line in lines:
                loc_id = str(line.location_id) if line.location else 'default'
                loc_name = line.location.name if line.location else cust.name
                if loc_id not in loc_map:
                    loc_map[loc_id] = {
                        "location_id": loc_id if loc_id != 'default' else None,
                        "location_name": loc_name,
                        "lines": []
                    }
                loc_map[loc_id]["lines"].append(line)

            for loc_id, loc_info in loc_map.items():
                loc_key = (str(cust.id), loc_info["location_id"] or 'default')
                if loc_key in seen_customer_locations:
                    continue
                seen_customer_locations.add(loc_key)

                row = {
                    "proposal_version_id": str(ver.id),
                    "proposal_id": str(prop.id),
                    "proposal_number": prop.proposal_number or "",
                    "client_id": str(cust.id),
                    "client_name": cust.name,
                    "location_id": loc_info["location_id"],
                    "location_name": loc_info["location_name"],
                    "overhead_per_guard": float(ver.overhead_per_guard if ver.overhead_per_guard is not None else 6000),
                    "service_charges_per_guard": float(ver.service_charges_per_guard if ver.service_charges_per_guard is not None else 3000),
                    "tax_wht_rate": float(ver.withholding_tax_rate if ver.withholding_tax_rate is not None else 7),
                    "sales_tax_rate": float(ver.tax_rate if ver.tax_rate is not None else 8),
                    "sales_tax_override": float(ver.sales_tax_override) if ver.sales_tax_override is not None else None,
                    "withholding_tax_override": float(ver.withholding_tax_override) if ver.withholding_tax_override is not None else None,
                    "sessi": float(ver.total_sessi if ver.total_sessi is not None else 0),
                    "eobi": float(ver.total_eobi if ver.total_eobi is not None else 0),

                    "sup_civ_qty": 0, "sup_civ_rate": 0, "sup_civ_sal": 0,
                    "sup_ex_qty": 0, "sup_ex_rate": 0, "sup_ex_sal": 0,
                    "sr_sup_civ_qty": 0, "sr_sup_civ_rate": 0, "sr_sup_civ_sal": 0,
                    "sr_sup_ex_qty": 0, "sr_sup_ex_rate": 0, "sr_sup_ex_sal": 0,
                    "guard_civ_qty": 0, "guard_civ_rate": 0, "guard_civ_sal": 0,
                    "guard_ex_qty": 0, "guard_ex_rate": 0, "guard_ex_sal": 0,
                    "hd_gd_civ_qty": 0, "hd_gd_civ_rate": 0, "hd_gd_civ_sal": 0,
                    "hd_gd_ex_qty": 0, "hd_gd_ex_rate": 0, "hd_gd_ex_sal": 0,
                    "cpo_civ_qty": 0, "cpo_civ_rate": 0, "cpo_civ_sal": 0,
                    "cpo_ex_qty": 0, "cpo_ex_rate": 0, "cpo_ex_sal": 0,
                    "cpo_ssg_qty": 0, "cpo_ssg_rate": 0, "cpo_ssg_sal": 0,
                    "lady_searcher_qty": 0, "lady_searcher_rate": 0, "lady_searcher_sal": 0,
                    "cctv_op_qty": 0, "cctv_op_rate": 0, "cctv_op_sal": 0,
                    "deo_qty": 0, "deo_rate": 0, "deo_sal": 0,
                    # Legacy fallback
                    "lady_cctv_qty": 0, "lady_cctv_rate": 0, "lady_cctv_sal": 0,
                }

                for l in loc_info["lines"]:
                    st_name = (l.service_type.name if l.service_type else "").lower()
                    st_code = (l.service_type.code if l.service_type else "").upper()
                    q = l.quantity or 0
                    r = float(l.client_rate or 0)
                    s = float(l.guard_salary or 0)

                    is_ex = ('ex' in st_name or 'arm' in st_name or '_EX' in st_code)
                    is_ssg = ('ssg' in st_name or 'commando' in st_name or 'SSG' in st_code)

                    if st_code == 'CPO_SSG' or (('cpo' in st_name or 'close protection' in st_name) and is_ssg):
                        row["cpo_ssg_qty"] += q; row["cpo_ssg_rate"] = r; row["cpo_ssg_sal"] = s
                    elif st_code == 'CPO_EX' or (('cpo' in st_name or 'close protection' in st_name) and is_ex):
                        row["cpo_ex_qty"] += q; row["cpo_ex_rate"] = r; row["cpo_ex_sal"] = s
                    elif st_code == 'CPO_CIV' or ('cpo' in st_name or 'close protection' in st_name):
                        row["cpo_civ_qty"] += q; row["cpo_civ_rate"] = r; row["cpo_civ_sal"] = s
                    elif st_code == 'SR_SUP_EX' or (('senior supervisor' in st_name or 'sr supervisor' in st_name) and is_ex):
                        row["sr_sup_ex_qty"] += q; row["sr_sup_ex_rate"] = r; row["sr_sup_ex_sal"] = s
                    elif st_code == 'SR_SUP_CIV' or ('senior supervisor' in st_name or 'sr supervisor' in st_name):
                        row["sr_sup_civ_qty"] += q; row["sr_sup_civ_rate"] = r; row["sr_sup_civ_sal"] = s
                    elif st_code == 'SUP_EX' or ('sup' in st_name and is_ex):
                        row["sup_ex_qty"] += q; row["sup_ex_rate"] = r; row["sup_ex_sal"] = s
                    elif st_code == 'SUP_CIV' or 'sup' in st_name:
                        row["sup_civ_qty"] += q; row["sup_civ_rate"] = r; row["sup_civ_sal"] = s
                    elif st_code == 'HD_GD_EX' or (('head' in st_name or 'senior guard' in st_name) and is_ex):
                        row["hd_gd_ex_qty"] += q; row["hd_gd_ex_rate"] = r; row["hd_gd_ex_sal"] = s
                    elif st_code == 'HD_GD_CIV' or ('head' in st_name or 'senior guard' in st_name):
                        row["hd_gd_civ_qty"] += q; row["hd_gd_civ_rate"] = r; row["hd_gd_civ_sal"] = s
                    elif st_code in ['CCTV_OP', 'CCTV'] or ('cctv' in st_name and 'lady' not in st_name):
                        row["cctv_op_qty"] += q; row["cctv_op_rate"] = r; row["cctv_op_sal"] = s
                    elif st_code in ['LADY_SEARCHER', 'LADY_CCTV'] or 'lady' in st_name or 'searcher' in st_name:
                        row["lady_searcher_qty"] += q; row["lady_searcher_rate"] = r; row["lady_searcher_sal"] = s
                        row["lady_cctv_qty"] += q; row["lady_cctv_rate"] = r; row["lady_cctv_sal"] = s
                    elif st_code == 'DEO' or 'deo' in st_name or 'data entry' in st_name:
                        row["deo_qty"] += q; row["deo_rate"] = r; row["deo_sal"] = s
                    elif st_code == 'GD_EX' or is_ex:
                        row["guard_ex_qty"] += q; row["guard_ex_rate"] = r; row["guard_ex_sal"] = s
                    else:
                        row["guard_civ_qty"] += q; row["guard_civ_rate"] = r; row["guard_civ_sal"] = s

                data.append(row)

        return Response(data, status=status.HTTP_200_OK)


class CostingGridBatchSyncView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        from erp_core.middleware import get_current_company
        from companies.models import Company
        from decimal import Decimal
        from django.db import transaction
        from crm.models import CRMEntity

        company_id = get_current_company() or request.META.get('HTTP_X_COMPANY_ID') or getattr(request.user, 'company_id', None)
        company = Company.objects.filter(id=company_id).first() if company_id else getattr(request.user, 'company', None)
        if not company and request.user.is_superuser:
            company = Company.objects.first()
        if not company:
            return Response({"error": "No company associated with user"}, status=status.HTTP_400_BAD_REQUEST)

        rows = request.data.get('rows', [])
        if not rows:
            return Response({"error": "No rows provided"}, status=status.HTTP_400_BAD_REQUEST)

        synced_count = 0

        role_definitions = [
            ("sup_civ", "Supervisor Civil", "SUP_CIV", "UNARMED"),
            ("sup_ex", "Supervisor Ex-Army", "SUP_EX", "UNARMED"),
            ("sr_sup_civ", "Senior Supervisor Civil", "SR_SUP_CIV", "UNARMED"),
            ("sr_sup_ex", "Senior Supervisor Ex-Army", "SR_SUP_EX", "UNARMED"),
            ("guard_civ", "Security Guard Civil", "GD_CIV", "UNARMED"),
            ("guard_ex", "Security Guard Ex-Army", "GD_EX", "UNARMED"),
            ("hd_gd_civ", "Head / Senior Guard Civil", "HD_GD_CIV", "UNARMED"),
            ("hd_gd_ex", "Head / Senior Guard Ex-Army", "HD_GD_EX", "UNARMED"),
            ("cpo_civ", "Close Protection Officer Civil", "CPO_CIV", "PISTOL"),
            ("cpo_ex", "Close Protection Officer Ex-Army", "CPO_EX", "PISTOL"),
            ("cpo_ssg", "Close Protection Officer Ex-SSG Commando", "CPO_SSG", "PISTOL"),
            ("lady_searcher", "Lady Searcher", "LADY_SEARCHER", "UNARMED"),
            ("cctv_op", "CCTV Operator", "CCTV_OP", "CCTV_OPERATOR"),
            ("deo", "Data Entry Operator (DEO)", "DEO", "UNARMED"),
        ]

        proposals_to_sync = set()
        try:
            with transaction.atomic():
                for row in rows:
                    client_name = str(row.get('client_name') or '').strip()
                    if not client_name:
                        continue

                    client_id = row.get('client_id')
                    customer = None
                    if client_id:
                        try:
                            customer = CRMEntity.objects.filter(id=client_id, company=company).first()
                        except Exception:
                            customer = None

                    if not customer:
                        customer = getattr(CRMEntity, 'all_objects', CRMEntity.objects).filter(company=company, name__iexact=client_name).first()
                        if customer and customer.is_deleted:
                            customer.is_deleted = False
                            customer.active = True
                            customer.save(update_fields=['is_deleted', 'active'])
                    if not customer:
                        customer = CRMEntity.objects.create(
                            company=company,
                            name=client_name,
                            entity_type='CUSTOMER',
                            active=True
                        )

                    loc_name = str(row.get('location_name') or '').strip() or client_name
                    location = getattr(ClientLocation, 'all_objects', ClientLocation.objects).filter(
                        company=company,
                        customer=customer,
                        name__iexact=loc_name
                    ).first()
                    if location and location.is_deleted:
                        location.is_deleted = False
                        location.is_active = True
                        location.save(update_fields=['is_deleted', 'is_active'])
                    if not location:
                        location = ClientLocation.objects.create(
                            company=company,
                            customer=customer,
                            name=loc_name,
                            is_active=True
                        )

                    proposal = getattr(SecurityProposal, 'all_objects', SecurityProposal.objects).filter(
                        company=company,
                        customer=customer
                    ).order_by('-created_at').first()
                    if proposal and proposal.is_deleted:
                        proposal.is_deleted = False
                        proposal.save(update_fields=['is_deleted'])
                    if not proposal:
                        proposal = SecurityProposal.objects.create(
                            company=company,
                            customer=customer,
                            title=f"{customer.name} - Final Requirements",
                            status=SecurityProposalStatus.ACTIVE,
                            is_handoff_ready=True
                        )
                    else:
                        if proposal.status != SecurityProposalStatus.ACTIVE or not proposal.is_handoff_ready:
                            proposal.status = SecurityProposalStatus.ACTIVE
                            proposal.is_handoff_ready = True
                            proposal.save(update_fields=['status', 'is_handoff_ready'])

                    def _parse_dec(v, default):
                        if v is None or str(v).strip() == '':
                            return Decimal(str(default))
                        try:
                            return Decimal(str(v))
                        except Exception:
                            return Decimal(str(default))

                    def _parse_opt_dec(v):
                        if v is None or str(v).strip() == '':
                            return None
                        try:
                            return Decimal(str(v))
                        except Exception:
                            return None

                    overhead = _parse_dec(row.get('overhead_per_guard'), 6000)
                    service_charges = _parse_dec(row.get('service_charges_per_guard'), 3000)
                    wht = _parse_dec(row.get('tax_wht_rate'), 7)
                    st_rate = _parse_dec(row.get('sales_tax_rate'), 8)
                    sales_tax_override = _parse_opt_dec(row.get('sales_tax_override'))
                    withholding_tax_override = _parse_opt_dec(row.get('withholding_tax_override'))
                    sessi = _parse_dec(row.get('sessi'), 0)
                    eobi = _parse_dec(row.get('eobi'), 0)

                    version = ProposalVersion.objects.filter(
                        company=company,
                        proposal=proposal
                    ).order_by('-version_number').first()
                    if not version:
                        version = ProposalVersion.objects.create(
                            company=company,
                            proposal=proposal,
                            version_number=1,
                            version_type='Final Requirement',
                            status='ACTIVE',
                            overhead_per_guard=overhead,
                            service_charges_per_guard=service_charges,
                            withholding_tax_rate=wht,
                            tax_rate=st_rate,
                            sales_tax_override=sales_tax_override,
                            withholding_tax_override=withholding_tax_override,
                            total_sessi=sessi,
                            total_eobi=eobi
                        )
                    else:
                        version.overhead_per_guard = overhead
                        version.service_charges_per_guard = service_charges
                        version.withholding_tax_rate = wht
                        version.tax_rate = st_rate
                        version.sales_tax_override = sales_tax_override
                        version.withholding_tax_override = withholding_tax_override
                        version.total_sessi = sessi
                        version.total_eobi = eobi
                        version.save()

                    existing_version_lines = list(ProposalServiceLine.objects.filter(
                        company=company,
                        proposal_version=version,
                        location=location
                    ).select_related('service_type'))

                    for prefix, role_name, role_code, default_weapon in role_definitions:
                        raw_qty = row.get(f"{prefix}_qty")
                        raw_rate = row.get(f"{prefix}_rate")
                        raw_sal = row.get(f"{prefix}_sal")

                        # Backwards compatibility fallback for lady_searcher / lady_cctv
                        if prefix == "lady_searcher" and (raw_qty is None or float(raw_qty or 0) == 0):
                            if row.get("lady_cctv_qty"):
                                raw_qty = row.get("lady_cctv_qty")
                                raw_rate = row.get("lady_cctv_rate")
                                raw_sal = row.get("lady_cctv_sal")

                        qty = int(float(raw_qty or 0))
                        rate = Decimal(str(raw_rate or 0))
                        sal = Decimal(str(raw_sal or 0))

                        st = SecurityServiceType.objects.filter(company=company, code=role_code).first()
                        if not st:
                            st = SecurityServiceType.objects.filter(company=company, name__iexact=role_name).first()
                        if not st:
                            st = SecurityServiceType.objects.create(
                                company=company,
                                name=role_name,
                                code=role_code,
                                is_active=True
                            )

                        # Match existing line by canonical role category
                        target_line = None
                        for cand in existing_version_lines:
                            cand_name = (cand.service_type.name or '').lower() if cand.service_type else ''
                            cand_code = (cand.service_type.code or '').upper() if cand.service_type else ''
                            is_cand_ex = ('ex' in cand_name or 'arm' in cand_name or '_EX' in cand_code)
                            is_cand_ssg = ('ssg' in cand_name or 'commando' in cand_name or 'SSG' in cand_code)

                            matched_role = None
                            if cand_code == 'CPO_SSG' or (('cpo' in cand_name or 'close protection' in cand_name) and is_cand_ssg):
                                matched_role = 'cpo_ssg'
                            elif cand_code == 'CPO_EX' or (('cpo' in cand_name or 'close protection' in cand_name) and is_cand_ex):
                                matched_role = 'cpo_ex'
                            elif cand_code == 'CPO_CIV' or ('cpo' in cand_name or 'close protection' in cand_name):
                                matched_role = 'cpo_civ'
                            elif cand_code == 'SR_SUP_EX' or (('senior supervisor' in cand_name or 'sr supervisor' in cand_name) and is_cand_ex):
                                matched_role = 'sr_sup_ex'
                            elif cand_code == 'SR_SUP_CIV' or ('senior supervisor' in cand_name or 'sr supervisor' in cand_name):
                                matched_role = 'sr_sup_civ'
                            elif cand_code == 'SUP_EX' or ('sup' in cand_name and is_cand_ex):
                                matched_role = 'sup_ex'
                            elif cand_code == 'SUP_CIV' or 'sup' in cand_name:
                                matched_role = 'sup_civ'
                            elif cand_code == 'HD_GD_EX' or (('head' in cand_name or 'senior guard' in cand_name) and is_cand_ex):
                                matched_role = 'hd_gd_ex'
                            elif cand_code == 'HD_GD_CIV' or ('head' in cand_name or 'senior guard' in cand_name):
                                matched_role = 'hd_gd_civ'
                            elif cand_code in ['CCTV_OP', 'CCTV'] or ('cctv' in cand_name and 'lady' not in cand_name):
                                matched_role = 'cctv_op'
                            elif cand_code in ['LADY_SEARCHER', 'LADY_CCTV'] or 'lady' in cand_name or 'searcher' in cand_name:
                                matched_role = 'lady_searcher'
                            elif cand_code == 'DEO' or 'deo' in cand_name or 'data entry' in cand_name:
                                matched_role = 'deo'
                            elif cand_code == 'GD_EX' or is_cand_ex:
                                matched_role = 'guard_ex'
                            else:
                                matched_role = 'guard_civ'

                            if matched_role == prefix:
                                if not target_line:
                                    target_line = cand
                                else:
                                    cand.delete()

                        if qty > 0:
                            if target_line:
                                target_line.service_type = st
                                target_line.quantity = qty
                                target_line.client_rate = rate
                                target_line.guard_salary = sal
                                target_line.save()
                            else:
                                ProposalServiceLine.objects.create(
                                    company=company,
                                    proposal_version=version,
                                    service_type=st,
                                    location=location,
                                    quantity=qty,
                                    client_rate=rate,
                                    guard_salary=sal,
                                    weapon_type=default_weapon,
                                    shift_hours='12_HOURS',
                                    billing_unit='MONTHLY'
                                )
                        else:
                            if target_line:
                                target_line.delete()

                    proposals_to_sync.add(proposal)
                    synced_count += 1

                for prop in proposals_to_sync:
                    try:
                        from security_crm.services.workflow import SecurityProposalWorkflowService
                        SecurityProposalWorkflowService.sync_proposal_and_locations_to_operations(prop)
                    except Exception as sync_err:
                        logger.warning(f"Error syncing proposal {prop.id} to operations: {sync_err}")

            return Response({
                "success": True,
                "synced_count": synced_count,
                "message": f"Successfully synced {synced_count} client location costing sheets into CRM and Operations!"
            }, status=status.HTTP_200_OK)

        except Exception as e:
            return Response({"error": f"Failed to sync costing grid: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)


class CostingGridImportExcelView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        import io
        import re
        import csv
        import openpyxl
        try:
            import xlrd
        except ImportError:
            xlrd = None

        from erp_core.middleware import get_current_company
        from companies.models import Company
        from decimal import Decimal
        from django.db import transaction
        from crm.models import CRMEntity
        from security_crm.models import (
            ClientLocation,
            SecurityProposal,
            ProposalVersion,
            ProposalServiceLine,
            SecurityServiceType
        )

        company_id = get_current_company() or request.META.get('HTTP_X_COMPANY_ID') or getattr(request.user, 'company_id', None)
        company = Company.objects.filter(id=company_id).first() if company_id else getattr(request.user, 'company', None)
        if not company and request.user.is_superuser:
            company = Company.objects.first()
        if not company:
            return Response({"error": "No company associated with user"}, status=status.HTTP_400_BAD_REQUEST)

        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            file_bytes = file_obj.read()
        except Exception as e:
            return Response({"error": f"Failed to read uploaded file: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)

        if not file_bytes:
            return Response({"error": "Uploaded file is empty"}, status=status.HTTP_400_BAD_REQUEST)

        max_rows = 0
        get_cell_val = None
        parse_errors = []

        # 1. Try openpyxl if file looks like a ZIP / OpenXML (.xlsx)
        if file_bytes.startswith(b'PK'):
            try:
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                target_sheet = wb.sheetnames[0]
                for s in wb.sheetnames:
                    s_lower = s.lower()
                    if 'sheet 1' in s_lower or 'sheet1' in s_lower or 'costing' in s_lower:
                        target_sheet = s
                        break
                ws = wb[target_sheet]
                max_rows = ws.max_row
                get_cell_val = lambda r, c: ws.cell(row=r, column=c).value
            except Exception as ex:
                parse_errors.append(f"openpyxl: {str(ex)}")

        # 2. Try xlrd for legacy binary Excel (.xls / BIFF8) or fallback
        if get_cell_val is None and xlrd is not None:
            try:
                wb_xls = xlrd.open_workbook(file_contents=file_bytes)
                sheet_names = wb_xls.sheet_names()
                target_sheet = sheet_names[0]
                for s in sheet_names:
                    s_lower = s.lower()
                    if 'sheet 1' in s_lower or 'sheet1' in s_lower or 'costing' in s_lower:
                        target_sheet = s
                        break
                sheet = wb_xls.sheet_by_name(target_sheet)
                max_rows = sheet.nrows
                get_cell_val = lambda r, c: sheet.cell_value(r - 1, c - 1) if (r - 1) < sheet.nrows and (c - 1) < sheet.ncols else None
            except Exception as ex:
                parse_errors.append(f"xlrd: {str(ex)}")

        # 3. Fallback: CSV / Delimited text
        if get_cell_val is None:
            try:
                text = file_bytes.decode('utf-8-sig', errors='ignore')
                sniffer = csv.Sniffer()
                dialect = sniffer.sniff(text[:2048])
                csv_rows = list(csv.reader(io.StringIO(text), dialect))
                max_rows = len(csv_rows)
                get_cell_val = lambda r, c: csv_rows[r - 1][c - 1] if (r - 1) < len(csv_rows) and (c - 1) < len(csv_rows[r - 1]) else None
            except Exception as ex:
                parse_errors.append(f"csv: {str(ex)}")

        if get_cell_val is None:
            err_msg = "; ".join(parse_errors) if parse_errors else "Unrecognized file format."
            return Response({"error": f"Failed to parse Excel file: {err_msg}"}, status=status.HTTP_400_BAD_REQUEST)

        role_specs = [
            ("Supervisor Ex-Army", "SUP_EX", 4, 5, 6, "UNARMED"),
            ("Supervisor Civil", "SUP_CIV", 7, 8, 9, "UNARMED"),
            ("Security Guard Ex-Army", "GD_EX", 10, 11, 12, "UNARMED"),
            ("Security Guard Civil", "GD_CIV", 13, 14, 15, "UNARMED"),
            ("Lady Searcher", "LADY_SEARCHER", 16, 17, 18, "UNARMED"),
            ("Close Protection Officer Ex-Army", "CPO_EX", 19, 20, 21, "PISTOL"),
            ("Close Protection Officer Civil", "CPO_CIV", 22, 23, 24, "PISTOL"),
        ]

        def parse_client_and_location(raw_name):
            raw = str(raw_name or '').replace('\xa0', ' ').strip()
            if not raw:
                return '', ''
            # Match pattern: Client Name (Location) or Client Name [Location]
            match = re.match(r'^(.*?)\s*[\(\[](.*?)[\)\]]\s*$', raw)
            if match:
                client = match.group(1).strip().strip(' -:')
                loc = match.group(2).strip().strip(' -:')
                if client and loc:
                    return client, loc
                elif client:
                    return client, client
                elif loc:
                    return loc, loc
            return raw, raw

        imported_clients_set = set()
        imported_locations = 0
        imported_lines = 0

        try:
            with transaction.atomic():
                for r in range(3, max_rows + 1):
                    loc_raw = get_cell_val(r, 2)
                    if not loc_raw:
                        continue
                    raw_str = str(loc_raw).replace('\xa0', ' ').strip()
                    if raw_str.lower() in ['total', 'totals', 'average', 'summary', 'location', 'client', '']:
                        continue

                    client_name, loc_name = parse_client_and_location(raw_str)
                    if not client_name:
                        continue

                    customer = getattr(CRMEntity, 'all_objects', CRMEntity.objects).filter(company=company, name__iexact=client_name).first()
                    if customer:
                        if customer.is_deleted:
                            customer.is_deleted = False
                            customer.active = True
                            customer.save(update_fields=['is_deleted', 'active'])
                    else:
                        customer = CRMEntity.objects.create(
                            company=company,
                            name=client_name,
                            entity_type='CUSTOMER',
                            active=True
                        )

                    location = getattr(ClientLocation, 'all_objects', ClientLocation.objects).filter(
                        company=company,
                        customer=customer,
                        name__iexact=loc_name
                    ).first()
                    if location:
                        if location.is_deleted:
                            location.is_deleted = False
                            location.is_active = True
                            location.save(update_fields=['is_deleted', 'is_active'])
                    else:
                        location = ClientLocation.objects.create(
                            company=company,
                            customer=customer,
                            name=loc_name,
                            is_active=True
                        )

                    proposal = getattr(SecurityProposal, 'all_objects', SecurityProposal.objects).filter(company=company, customer=customer).order_by('-created_at').first()
                    if proposal and proposal.is_deleted:
                        proposal.is_deleted = False
                        proposal.save(update_fields=['is_deleted'])
                    if not proposal:
                        proposal = SecurityProposal.objects.create(
                            company=company,
                            customer=customer,
                            title=f"{customer.name} - Commercial Proposal",
                            status='DRAFT'
                        )

                    sessi_val = get_cell_val(r, 31)
                    eobi_val = get_cell_val(r, 32)
                    try:
                        sessi = Decimal(str(sessi_val or 0)) if sessi_val else Decimal('0.00')
                    except Exception:
                        sessi = Decimal('0.00')
                    try:
                        eobi = Decimal(str(eobi_val or 0)) if eobi_val else Decimal('0.00')
                    except Exception:
                        eobi = Decimal('0.00')

                    version = ProposalVersion.objects.filter(company=company, proposal=proposal, is_frozen=False).order_by('-version_number').first()
                    if not version:
                        version = ProposalVersion.objects.create(
                            company=company,
                            proposal=proposal,
                            version_number=1,
                            version_type='Initial Proposal',
                            status='DRAFT',
                            overhead_per_guard=Decimal('6000.00'),
                            service_charges_per_guard=Decimal('3000.00'),
                            total_sessi=sessi,
                            total_eobi=eobi
                        )
                    else:
                        dirty_fields = []
                        if sessi > 0 and version.total_sessi != sessi:
                            version.total_sessi = sessi
                            dirty_fields.append('total_sessi')
                        if eobi > 0 and version.total_eobi != eobi:
                            version.total_eobi = eobi
                            dirty_fields.append('total_eobi')
                        if dirty_fields:
                            version.save(update_fields=dirty_fields)

                    # Ensure idempotency by replacing any existing lines for this location in this proposal version
                    ProposalServiceLine.objects.filter(
                        company=company,
                        proposal_version=version,
                        location=location
                    ).delete()

                    location_has_lines = False
                    for role_title, role_code, rate_col, sal_col, qty_col, weapon in role_specs:
                        try:
                            qty_val = get_cell_val(r, qty_col)
                            if not qty_val:
                                continue
                            qty = int(float(qty_val))
                            if qty <= 0:
                                continue

                            rate_val = get_cell_val(r, rate_col) or 0
                            sal_val = get_cell_val(r, sal_col) or 0
                            rate = Decimal(str(rate_val))
                            salary = Decimal(str(sal_val))

                            st = SecurityServiceType.objects.filter(company=company, code=role_code).first()
                            if not st:
                                st = SecurityServiceType.objects.filter(company=company, name__iexact=role_title).first()
                            if not st:
                                st = SecurityServiceType.objects.create(
                                    company=company,
                                    name=role_title,
                                    code=role_code,
                                    is_active=True
                                )

                            ProposalServiceLine.objects.create(
                                company=company,
                                proposal_version=version,
                                service_type=st,
                                location=location,
                                quantity=qty,
                                client_rate=rate,
                                guard_salary=salary,
                                weapon_type=weapon,
                                shift_hours='12_HOURS',
                                billing_unit='MONTHLY'
                            )
                            imported_lines += 1
                            location_has_lines = True
                        except Exception:
                            continue

                    if location_has_lines:
                        imported_locations += 1
                        imported_clients_set.add(customer.id)

            imported_clients = len(imported_clients_set)
            return Response({
                "success": True,
                "imported_clients": imported_clients,
                "imported_locations": imported_locations,
                "imported_lines": imported_lines,
                "message": f"Successfully imported {imported_clients} clients across {imported_locations} locations with {imported_lines} staffing lines from Sheet 1!"
            }, status=status.HTTP_200_OK)

        except Exception as e:
            return Response({"error": f"Failed to import Excel file: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)


