import os
import openpyxl
from datetime import time as dtime
from decimal import Decimal
from django.conf import settings
from django.db import transaction
from hrm.models import Employee, Department, Designation
from ..models import (
    TechnicianProfile, MobileMonthlyAttendance, MobileMonthlyPayroll,
    MobileDepartment
)
from .payroll_service import MobilePayrollService

class MobileExcelSeedService:
    @staticmethod
    @transaction.atomic
    def seed_softech_sample_data(company):
        """
        Reads 'Softech Data/HRM - (September 2026).xlsx' and seeds:
        1. Universal Departments and Employees
        2. TechnicianProfile with phones, IMEIs, shift times, and OT rates
        3. September 2026 MobileMonthlyAttendance grid
        4. September 2026 MobileMonthlyPayroll records
        """
        excel_path = os.path.join(settings.BASE_DIR, 'Softech Data', 'HRM - (September 2026).xlsx')
        if not os.path.exists(excel_path):
            return {'status': 'error', 'message': f"File not found: {excel_path}"}

        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws_roster = wb['HRM (Overview)']
        ws_data = wb['Data Sheet']

        # Ensure departments exist
        dept_map = {
            'hardware': MobileDepartment.HARDWARE,
            'software': MobileDepartment.SOFTWARE,
            'panel-glass': MobileDepartment.PANEL_GLASS,
            'sales': MobileDepartment.SALES,
            'w.o': MobileDepartment.SALES,
            'management': MobileDepartment.MANAGEMENT,
        }

        created_employees = {}
        for r in range(5, ws_roster.max_row + 1):
            emp_code = str(ws_roster.cell(r, 1).value or '').strip()
            if not emp_code or not emp_code.upper().startswith('EMP'):
                continue

            raw_dept = str(ws_roster.cell(r, 2).value or '').strip().lower()
            dept_choice = dept_map.get(raw_dept, MobileDepartment.HARDWARE)
            full_name = str(ws_roster.cell(r, 3).value or '').strip()
            if not full_name:
                continue

            # Parse shift time
            shift_val = ws_roster.cell(r, 4).value
            shift_time = dtime(12, 0)
            if isinstance(shift_val, dtime):
                shift_time = shift_val
            elif isinstance(shift_val, str) and ':' in shift_val:
                parts = shift_val.split(':')
                try:
                    shift_time = dtime(int(parts[0]), int(parts[1]))
                except Exception:
                    pass

            half_day_limit = int(ws_roster.cell(r, 5).value or 0)
            med_limit = int(ws_roster.cell(r, 6).value or 1)
            
            ot_amount_raw = ws_roster.cell(r, 9).value
            try:
                ot_rate = Decimal(str(ot_amount_raw))
            except Exception:
                ot_rate = Decimal('250.00')

            phone_1 = str(ws_roster.cell(r, 10).value or '').strip().replace('-', '')
            phone_2 = str(ws_roster.cell(r, 11).value or '').strip().replace('-', '')
            imei_1 = str(ws_roster.cell(r, 12).value or '').strip().replace('-', '')
            imei_2 = str(ws_roster.cell(r, 13).value or '').strip().replace('-', '')

            # Create or update universal Employee
            dept_obj, _ = Department.objects.get_or_create(
                company=company,
                name=dept_choice.capitalize()
            )

            name_parts = full_name.split(' ', 1)
            first_name = name_parts[0]
            last_name = name_parts[1] if len(name_parts) > 1 else ''

            emp, _ = Employee.objects.update_or_create(
                company=company,
                employee_code=emp_code,
                defaults={
                    'first_name': first_name,
                    'last_name': last_name,
                    'department': dept_obj,
                    'is_active': True,
                    'classification': 'DIRECT',
                    'employment_status': 'ACTIVE'
                }
            )

            # Create or update TechnicianProfile
            TechnicianProfile.objects.update_or_create(
                company=company,
                employee=emp,
                defaults={
                    'department_type': dept_choice,
                    'shift_time': shift_time,
                    'allowed_half_days': half_day_limit,
                    'allowed_medical_half_days': med_limit,
                    'ot_hourly_rate': ot_rate,
                    'phone_1_model': phone_1,
                    'phone_1_imei_1': imei_1,
                    'phone_2_model': phone_2,
                    'phone_2_imei_1': imei_2,
                }
            )
            created_employees[emp_code] = emp

        # Now read September 2026 attendance from 'Data Sheet'
        for r in range(9, 25):
            code_cell = str(ws_data.cell(r, 2).value or '').strip()
            if not code_cell or code_cell not in created_employees:
                continue

            emp = created_employees[code_cell]
            days_data = {}
            # Days 1 to 30 are in cols 5 to 34
            for day_num in range(1, 31):
                col_idx = 4 + day_num
                cell_val = ws_data.cell(r, col_idx).value
                if cell_val:
                    days_data[str(day_num)] = str(cell_val).strip().upper()

            att, _ = MobileMonthlyAttendance.objects.update_or_create(
                company=company,
                employee=emp,
                year=2026,
                month=9,
                defaults={'days_data': days_data}
            )
            att.recalculate_counters()
            att.save()

            # Create initial Payroll Record
            MobilePayrollService.save_or_update_payroll(
                company=company,
                employee=emp,
                year=2026,
                month=9,
                data={
                    'basic_salary': Decimal('20000.00') if emp.employee_code != 'EMP25' else Decimal('18000.00'),
                    'loan_principal': Decimal('2990.00') if emp.employee_code == 'EMP10' else Decimal('0.00'),
                    'loan_deduction': Decimal('0.00'),
                    'damaged_parts': Decimal('1900.00') if emp.employee_code == 'EMP43' else Decimal('0.00'),
                    'bill_deduction': Decimal('200.00') if emp.employee_code == 'EMP43' else Decimal('0.00'),
                }
            )

        # Seed Company Policy
        from ..models import CompanyPolicy, OnboardingCandidate, DEFAULT_SOFTECH_POLICY
        from datetime import date
        CompanyPolicy.objects.get_or_create(
            company=company,
            title='Rules & Regulations',
            defaults={
                'content': DEFAULT_SOFTECH_POLICY,
                'version': '1.0',
                'is_active': True,
                'created_by_name': 'Sheikh Muhammad Afnan (CEO)'
            }
        )

        # Seed Sample Onboarding Candidates (from Softech Appointment Letters)
        OnboardingCandidate.objects.get_or_create(
            company=company,
            full_name='Shahzaib Akbar',
            defaults={
                'phone': '0300-1234567',
                'address': '1st floor Shop#169 Star City Mall, Saddar, Karachi.',
                'position': 'W.O (Management)',
                'department_type': MobileDepartment.MANAGEMENT,
                'offered_salary': Decimal('18000.00'),
                'allowance': '-',
                'backup_work': 'W.O (Management)',
                'working_hours': '11:00 to 9:00 (Closing)',
                'notice_period': 'One Month',
                'appointment_date': date(2025, 6, 11),
                'appointment_month': 'June',
                'ceo_name': 'Sheikh Muhammad Afnan',
                'director_name': 'Shahzaib Ali',
                'status': 'OFFERED',
                'first_week_checklist': {
                    'induction': True,
                    'toolset_assigned': True,
                    'policy_signed': True,
                    'login_created': False
                }
            }
        )

        OnboardingCandidate.objects.get_or_create(
            company=company,
            full_name='Subhan Shaikh',
            defaults={
                'phone': '0321-9876543',
                'address': '1st floor Shop#169 Star City Mall, Saddar, Karachi.',
                'position': 'Mobile Hardware Technician',
                'department_type': MobileDepartment.HARDWARE,
                'offered_salary': Decimal('20000.00'),
                'allowance': 'Travel 1000 Pkr/= (Every Week)',
                'backup_work': '-',
                'working_hours': '11:00 to 9:00',
                'notice_period': 'One – Two Month',
                'appointment_date': date(2025, 10, 1),
                'appointment_month': 'October',
                'ceo_name': 'Sheikh Muhammad Afnan',
                'director_name': 'Shahzaib Ali',
                'status': 'OFFERED',
                'first_week_checklist': {
                    'induction': True,
                    'toolset_assigned': True,
                    'policy_signed': True,
                    'login_created': True
                }
            }
        )

        return {
            'status': 'success',
            'seeded_employees': len(created_employees),
            'message': f"Successfully seeded {len(created_employees)} technicians, policies, onboarding candidates, and September 2026 attendance/payroll."
        }
